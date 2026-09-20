"""Run:  python3 -m pytest -q

These use hand-written emails and documents that are NOT in the organisers'
dataset — different wording, labels, languages and layouts — to check the
system generalises instead of memorising the sample data.
"""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import pytest  # noqa: E402

from app import llm  # noqa: E402
from app.classify import classify  # noqa: E402
from app.fields import (fields_from_pairs, is_placeholder, label_to_field, norm_party, norm_port,  # noqa: E402
                        parse_container_count, parse_weight_kg)
from app.parsing import parse_attachment  # noqa: E402
from app.pipeline import process_email, to_submission  # noqa: E402


@pytest.fixture(autouse=True)
def no_llm(monkeypatch):
    monkeypatch.setattr(llm, "available", lambda: False)
    monkeypatch.setattr(llm, "can_read_pdf", lambda: False)


# ------------------------------------------------------------ normalisers

@pytest.mark.parametrize("a,b", [
    ("MOORIM SP CO., LTD", "Moorim SP Co Ltd"),
    ("KPP-ANTALIS (SINGAPORE) PTE. LTD.", "KPP ANTALIS (SINGAPORE) PTE LTD"),
    ("BALL & DOGGETT AUSTRALIA PTY LTD\n43-45 METROPOLITAN ROAD", "BALL AND DOGGETT AUSTRALIA PTY LTD"),
    ("ACME TRADING COMPANY LIMITED", "ACME TRADING CO LTD"),
])
def test_same_party_different_formatting(a, b):
    assert norm_party(a) == norm_party(b)


def test_different_companies_stay_different():
    # real, distinct entities that differ only by a suffix — must NOT match
    assert norm_party("APRIL FINE PAPER TRADING") != norm_party("APRIL FINE PAPER TRADING (MIDDLE EAST) FZE")


@pytest.mark.parametrize("a,b,same", [
    ("PORT KLANG (WESTPORT), MALAYSIA (MYPKG)", "Port Klang, Malaysia", True),
    ("HOCHIMINH CITY, VIETNAM (VNSGN)", "Ho Chi Minh City", True),
    ("MOMBASA, KENYA (KEMBA)", "TUTICORIN, INDIA (KEMBA)", False),   # same code, different city
])
def test_ports(a, b, same):
    assert (norm_port(a) == norm_port(b)) is same


@pytest.mark.parametrize("text,n", [("3 x 40'HC", 3), ("2x40HC + 1x20GP", 3), ("12 x 20'FCL", 12), ("5", 5)])
def test_container_count(text, n):
    assert parse_container_count(text) == n


@pytest.mark.parametrize("text,kg", [("21,577 KG", 21577), ("243588", 243588), ("21.5 MT", 21500), (131322, 131322)])
def test_weight(text, kg):
    assert parse_weight_kg(text) == kg


@pytest.mark.parametrize("v", ["???", "____MT", "TBA", "", None, "_______ MTS"])
def test_placeholders(v):
    assert is_placeholder(v)


# ------------------------------------------------------------ unseen layouts

SI_MALAY = """ARAHAN PENGHANTARAN (SHIPPING INSTRUCTION)
Pengirim: ACME PAPER SDN BHD
  12 JALAN INDUSTRI, SHAH ALAM
Penerima: GLOBAL BOOKS LLC
Pihak Dimaklumkan: GLOBAL BOOKS LLC
Pelabuhan Muat: PORT KLANG, MALAYSIA
Pelabuhan Bongkar: JEBEL ALI, UAE
Jumlah Kontainer: 2x40HC + 1x20GP
Berat Kasar: 55.2 MT
"""

BL_ENGLISH = """BILL OF LADING (DRAFT)
Shipper: Acme Paper Sdn. Bhd.
Consignee: GLOBAL BOOKS L.L.C.
Notify Party: GLOBAL BOOKS LLC
Load Port: Port Klang (Westport), Malaysia (MYPKG)
Discharge Port: Jebel Ali, UAE
No. of Containers: 3
Gross Weight (KG): 55,200 KG
"""


def test_malay_labels_are_understood():
    doc = parse_attachment("si.txt", SI_MALAY.encode())
    assert doc.doc_type == "SI"
    f = fields_from_pairs(doc.pairs)
    assert set(f) == {"shipper", "consignee", "notify_party", "port_of_loading",
                      "port_of_discharge", "container_count", "gross_weight_kg"}


@pytest.mark.parametrize("label,field", [
    ("Pihak untuk dihubungi", "notify_party"),
    ("Penerima kiriman", "consignee"),
    ("Pihak penghantar", "shipper"),
    ("Pelabuhan pemuatan", "port_of_loading"),
    ("Pelabuhan destinasi", "port_of_discharge"),
    ("Kuantiti kontena", "container_count"),
    ("Jumlah berat kasar", "gross_weight_kg"),
])
def test_additional_malay_field_labels(label, field):
    assert label_to_field(label) == field


@pytest.mark.parametrize("label,field", [
    ("通知方", "notify_party"),
    ("提货人", "consignee"),
    ("托运人", "shipper"),
    ("起运港", "port_of_loading"),
    ("目的港", "port_of_discharge"),
    ("集装箱数量", "container_count"),
    ("总毛重", "gross_weight_kg"),
])
def test_additional_chinese_field_labels(label, field):
    assert label_to_field(label) == field


def _email(body, atts, subject="Please check draft BL"):
    return {"email_id": "t1", "from": "a@b.com", "subject": subject, "body": body, "attachments": list(atts)}


def test_end_to_end_formatting_only_is_ok():
    files = {"si.txt": SI_MALAY.encode(), "bl.txt": BL_ENGLISH.encode()}
    e = _email("Attached SI and draft BL, please check the BL against the SI.", files)
    r = process_email(e, files.__getitem__, use_llm=False)
    assert r["category"] == "BL_COMPARISON"
    assert r["status"] == "OK", r["comparison"]


def test_end_to_end_catches_the_one_wrong_field():
    bad = BL_ENGLISH.replace("Discharge Port: Jebel Ali, UAE", "Discharge Port: Dammam, Saudi Arabia")
    files = {"si.txt": SI_MALAY.encode(), "bl.txt": bad.encode()}
    r = process_email(_email("Please verify the draft BL matches the SI.", files), files.__getitem__, use_llm=False)
    assert r["status"] == "MISMATCH"
    assert r["defect_fields"] == ["port_of_discharge"]
    sub = to_submission(r)
    assert sub["has_defect"] and sub["defect_fields"] == ["port_of_discharge"]


def test_blank_si_value_goes_to_a_person_not_a_mismatch():
    si = SI_MALAY.replace("Berat Kasar: 55.2 MT", "Berat Kasar: TBA")
    files = {"si.txt": si.encode(), "bl.txt": BL_ENGLISH.encode()}
    r = process_email(_email("Please check the draft BL against the SI.", files), files.__getitem__, use_llm=False)
    assert r["status"] == "NEEDS_REVIEW" and r["review_reason"] == "missing_value"


def test_wrong_document_goes_to_a_person():
    pl = "PACKING LIST\nShipper: ACME PAPER SDN BHD\nCarton No. 1\n"
    files = {"si.txt": SI_MALAY.encode(), "bl.txt": pl.encode()}
    r = process_email(_email("Please check the draft BL against the SI.", files), files.__getitem__, use_llm=False)
    assert r["status"] == "NEEDS_REVIEW" and r["review_reason"] == "wrong_doc_type"


def test_empty_and_broken_files_go_to_a_person():
    files = {"si.txt": SI_MALAY.encode(), "bl.pdf": b"%PDF-1.5 garbage"}
    r = process_email(_email("Please check the draft BL against the SI.", files), files.__getitem__, use_llm=False)
    assert r["status"] == "NEEDS_REVIEW" and r["review_reason"] == "unreadable"
    files = {"si.txt": SI_MALAY.encode(), "bl.txt": b""}
    r = process_email(_email("Please check the draft BL against the SI.", files), files.__getitem__, use_llm=False)
    assert r["review_reason"] == "unreadable"


def test_binary_formats():
    import docx
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in [("BL INSTRUCTION", ""), ("Shipper/Exporter", "ACME PAPER SDN BHD | 12 JALAN"), ("Consignee", "GLOBAL BOOKS LLC"),
                ("Notify Party", "GLOBAL BOOKS LLC"), ("POL", "PORT KLANG"), ("POD", "JEBEL ALI"),
                ("Container Count", "3 x 40'HC"), ("GROSS WEIGHT", 55200)]:
        ws.append(row)
    xb = io.BytesIO(); wb.save(xb)
    d = docx.Document(); d.add_paragraph("BILL OF LADING (DRAFT)")
    t = d.add_table(rows=0, cols=2)
    for k, v in [("Shipper (发货人)", "ACME PAPER SDN BHD"), ("Consignee (收货人)", "GLOBAL BOOKS LLC"), ("Notify (通知人)", "GLOBAL BOOKS LLC"),
                 ("PORT OF LOADING (装货港)", "PORT KLANG"), ("POD (卸货港)", "JEBEL ALI"), ("Total Containers (箱数)", "3 x 40'HC"),
                 ("Gross Wt (kgs) (毛重 KGS)", "55,200")]:
        c = t.add_row().cells; c[0].text, c[1].text = k, v
    db = io.BytesIO(); d.save(db)
    files = {"si.xlsx": xb.getvalue(), "bl.docx": db.getvalue()}
    r = process_email(_email("Please check the draft BL against the SI.", files), files.__getitem__, use_llm=False)
    assert r["status"] == "OK", r["comparison"]


# ------------------------------------------------------------ triage on new wording

@pytest.mark.parametrize("body,subject,cat", [
    ("Dilampirkan SI dan draf BL. Sila semak draf BL berbanding SI.", "Semakan draf BL", "BL_COMPARISON"),
    ("附件是SI和提单草稿，请核对提单是否与SI一致。", "提单草稿核对", "BL_COMPARISON"),
    ("Can you look over the attached bill of lading and make sure it lines up with our shipping instruction?", "BL", "BL_COMPARISON"),
    ("We were charged twice for port handling. Please send a corrected bill.", "Question", "INVOICE_QUERY"),
    ("Here are the shipping instructions for our next order to Jebel Ali.", "New shipment", "SI_REQUEST"),
    ("Congratulations! You have won a free container. Claim your prize now.", "Lucky shipper", "SPAM"),
])
def test_triage_new_wording(body, subject, cat):
    r = classify({"from": "x@y.com", "subject": subject, "body": body, "attachments": []}, use_llm=False)
    assert r["category"] == cat, r["scores"]


@pytest.mark.parametrize("body", [
    "Sila sahkan draf bil muatan untuk penghantaran ini.",
    "Mohon periksa konosemen yang dilampirkan.",
    "Tolong bandingkan bil muatan dengan arahan penghantaran.",
    "Pastikan BL sepadan dengan SI sebelum dihantar.",
    "Ada percanggahan antara konosemen dan arahan penghantaran; sila teliti.",
])
def test_additional_malay_bl_phrases(body):
    r = classify({"from": "ops@example.my", "subject": "Dokumen penghantaran", "body": body,
                  "attachments": ["si.txt", "bl.txt"]}, use_llm=False)
    assert r["category"] == "BL_COMPARISON", r["scores"]


@pytest.mark.parametrize("body", [
    "请确认提单草稿。",
    "请审核附件中的提单。",
    "请比较提单和装运指示。",
    "请检查提单与托运指示是否相符。",
    "提单和装运指示之间有差异，请核对。",
])
def test_additional_chinese_bl_phrases(body):
    r = classify({"from": "ops@example.cn", "subject": "运输文件", "body": body,
                  "attachments": ["si.txt", "bl.txt"]}, use_llm=False)
    assert r["category"] == "BL_COMPARISON", r["scores"]


# ------------------------------------------------------------ DeepSeek wiring (no network)

def test_deepseek_request_shape(monkeypatch, tmp_path):
    import importlib
    monkeypatch.undo()                          # use the real llm.available()
    monkeypatch.setenv("LLM_PROVIDER", "deepseek")
    monkeypatch.setenv("DEEPSEEK_API_KEY", "test-key")
    # pin the model here: without this the test reads the developer's own .env and
    # fails the moment the team switches models (deepseek-chat -> deepseek-flash).
    monkeypatch.setenv("DEEPSEEK_MODEL", "deepseek-chat")
    monkeypatch.setenv("LLM_CACHE_DIR", str(tmp_path))
    mod = importlib.reload(llm)
    sent = {}

    class Resp:
        status_code = 200
        def raise_for_status(self): pass
        def json(self): return {"choices": [{"message": {"content": '{"category": "SPAM", "confidence": 0.9, "reason": "prize scam"}'}}]}

    def fake_post(url, headers, json, timeout):
        sent.update(url=url, headers=headers, body=json)
        return Resp()

    monkeypatch.setattr(mod.requests, "post", fake_post)
    monkeypatch.setattr(mod, "MIN_INTERVAL", 0)
    out = mod.classify_email({"from": "a", "subject": "b", "body": "c"}, {"SPAM": "spam"})
    assert out["category"] == "SPAM"
    assert sent["url"] == "https://api.deepseek.com/chat/completions"
    assert sent["headers"]["Authorization"] == "Bearer test-key"
    assert sent["body"]["model"] == "deepseek-chat"
    assert sent["body"]["response_format"] == {"type": "json_object"}
    assert "json" in sent["body"]["messages"][0]["content"].lower()   # DeepSeek JSON mode requires it
    assert not mod.can_read_pdf()                                        # scans → a person
    importlib.reload(llm)


def test_scans_go_to_gemini_even_when_deepseek_is_main(monkeypatch, tmp_path):
    import importlib
    monkeypatch.undo()
    monkeypatch.setenv("LLM_PROVIDER", "deepseek")
    monkeypatch.setenv("DEEPSEEK_API_KEY", "ds-key")
    monkeypatch.setenv("GEMINI_API_KEY", "gm-key")
    monkeypatch.setenv("LLM_CACHE_DIR", str(tmp_path))
    mod = importlib.reload(llm)
    urls = []

    class Resp:
        status_code = 200
        def __init__(self, url): self.url = url
        def raise_for_status(self): pass
        def json(self):
            if "googleapis" in self.url:
                return {"candidates": [{"content": {"parts": [{"text": '{"doc_type": "BL", "fields": {}, "readable": true}'}]}}]}
            return {"choices": [{"message": {"content": '{"category": "SPAM"}'}}]}

    def fake_post(url, headers, json, timeout):
        urls.append(url)
        return Resp(url)

    monkeypatch.setattr(mod.requests, "post", fake_post)
    monkeypatch.setattr(mod, "MIN_INTERVAL", 0)
    assert mod.provider() == "deepseek" and mod.can_read_pdf()
    mod.extract_fields(None, pdf_bytes=b"%PDF scan")                      # scan → Gemini
    mod.classify_email({"from": "a", "subject": "b", "body": "c"}, {"SPAM": "x"})  # text → DeepSeek
    assert "googleapis.com" in urls[0] and urls[1] == "https://api.deepseek.com/chat/completions"
    importlib.reload(llm)


# ------------------------------------------------------------ Supabase key handling

def test_supabase_new_style_key_is_not_sent_as_bearer():
    from app.store import SupabaseStore
    new = SupabaseStore("https://x.supabase.co", "sb_secret_abc")
    assert new.h["apikey"] == "sb_secret_abc" and "Authorization" not in new.h
    legacy = SupabaseStore("https://x.supabase.co", "eyJhbGciOi.payload.sig")
    assert legacy.h["Authorization"] == "Bearer eyJhbGciOi.payload.sig"


# ------------------------------------- AI second opinion (advisory only)

def _mismatching_pair():
    bad = BL_ENGLISH.replace("Discharge Port: Jebel Ali, UAE", "Discharge Port: Dammam, Saudi Arabia")
    return {"si.txt": SI_MALAY.encode(), "bl.txt": bad.encode()}


@pytest.mark.parametrize("agrees", [True, False])
def test_second_opinion_never_changes_the_submitted_answer(monkeypatch, agrees):
    """The AI reviewer is a second pair of eyes, not a decision maker.

    Whatever it says — even flatly contradicting the rules — the category, status
    and defect fields we submit must be byte-identical to the run without it.
    """
    files = _mismatching_pair()
    email = _email("Please verify the draft BL matches the SI.", files)

    baseline = to_submission(process_email(email, files.__getitem__, use_llm=False))

    monkeypatch.setattr(llm, "available", lambda: True)
    monkeypatch.setattr(llm, "extract_fields", lambda *a, **k: None)
    monkeypatch.setattr(llm, "second_opinion",
                        lambda *a, **k: {"agrees": agrees, "note": "checked", "confidence": 0.9, "model": "test"})
    r = process_email(email, files.__getitem__, use_llm=True, second_opinion=True)

    assert to_submission(r) == baseline
    assert r["second_opinion"]["agrees"] is agrees
    assert "double-checked the mismatch as a second reviewer" in r["ai_used"]


def test_second_opinion_is_off_by_default():
    files = _mismatching_pair()
    r = process_email(_email("Please verify the draft BL matches the SI.", files), files.__getitem__, use_llm=False)
    assert "second_opinion" not in r


def test_second_opinion_failure_is_not_fatal(monkeypatch):
    """A dead AI key mid-demo must leave the decision untouched, not crash the email."""
    files = _mismatching_pair()
    monkeypatch.setattr(llm, "available", lambda: True)
    monkeypatch.setattr(llm, "extract_fields", lambda *a, **k: None)
    monkeypatch.setattr(llm, "second_opinion", lambda *a, **k: None)
    r = process_email(_email("Please verify the draft BL matches the SI.", files),
                      files.__getitem__, use_llm=True, second_opinion=True)
    assert r["status"] == "MISMATCH" and "second_opinion" not in r


def test_ai_used_is_empty_when_rules_do_everything():
    files = {"si.txt": SI_MALAY.encode(), "bl.txt": BL_ENGLISH.encode()}
    r = process_email(_email("Attached SI and draft BL.", files), files.__getitem__, use_llm=False)
    assert r["status"] == "OK" and r["ai_used"] == []
